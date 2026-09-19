from io import BytesIO
import json
from unittest.mock import patch

from django.core.files.uploadedfile import SimpleUploadedFile
from django.test import TestCase, override_settings
from django.urls import reverse
from rest_framework import status
from rest_framework.test import APITestCase
from docx import Document
from pypdf import PdfWriter
from apps.accounts.models import User
from apps.resumes.models import Resume

from .models import ResumeAnalysis
from .services.analysis_service import (
    analyze_resume,
    extract_and_save_resume_text,
)
from .services.gemini_service import (
    GeminiAPIError,
    InvalidGeminiResponseError,
    analyze_resume_with_gemini,
    validate_resume_analysis_result,
)
from .services.text_extractor import (
    clean_text,
    extract_resume_text,
)

VALID_GEMINI_RESULT = {
    "resume_score": 78,
    "candidate_summary": "Computer science student with Python and Java skills.",
    "skills": ["Python", "Java", "SQL"],
    "education": ["Computer Science Engineering"],
    "experience": [],
    "projects": [],
    "certifications": [],
    "strengths": ["Strong programming fundamentals"],
    "missing_skills": ["Cloud platforms"],
    "improvement_suggestions": ["Add project descriptions with measurable outcomes"],
    "recommended_roles": ["Software Engineer", "Backend Developer"],
}


class ResumeTextExtractionTests(TestCase):

    def create_pdf(self):
        """
        Create a minimal PDF file for testing.
        """

        writer = PdfWriter()

        writer.add_blank_page(
            width=612,
            height=792,
        )

        pdf_buffer = BytesIO()
        writer.write(pdf_buffer)

        pdf_buffer.seek(0)

        return SimpleUploadedFile(
            "test_resume.pdf",
            pdf_buffer.read(),
            content_type="application/pdf",
        )

    def create_docx(self):
        """
        Create a DOCX file containing test resume text.
        """

        document = Document()

        document.add_paragraph("Shifana Barveen")
        document.add_paragraph("Python Java SQL")
        document.add_paragraph("Computer Science Engineering")

        docx_buffer = BytesIO()
        document.save(docx_buffer)

        docx_buffer.seek(0)

        return SimpleUploadedFile(
            "test_resume.docx",
            docx_buffer.read(),
            content_type=(
                "application/vnd.openxmlformats-officedocument"
                ".wordprocessingml.document"
            ),
        )

    def test_clean_text(self):
        text = """
        Python     Java


        SQL
        """

        result = clean_text(text)

        self.assertEqual(
            result,
            "Python Java\nSQL",
        )

    def test_pdf_extraction(self):
        pdf_file = self.create_pdf()

        result = extract_resume_text(pdf_file)

        self.assertIsInstance(result, str)

    def test_docx_extraction(self):
        docx_file = self.create_docx()

        result = extract_resume_text(docx_file)

        self.assertIn("Shifana Barveen", result)
        self.assertIn("Python Java SQL", result)
        self.assertIn(
            "Computer Science Engineering",
            result,
        )

    def test_unsupported_file_type(self):
        txt_file = SimpleUploadedFile(
            "resume.txt",
            b"This is a text file.",
            content_type="text/plain",
        )

        with self.assertRaises(ValueError):
            extract_resume_text(txt_file)

    def test_extract_and_save_resume_text(self):
        user = User.objects.create_user(
            email="student@example.com",
            password="TestPassword123",
        )

        docx_file = self.create_docx()

        resume = Resume.objects.create(
            student=user,
            file=docx_file,
            original_filename="test_resume.docx",
            is_active=True,
        )

        analysis = extract_and_save_resume_text(resume)

        self.assertIsNotNone(analysis)
        self.assertEqual(
            analysis.resume,
            resume,
        )
        self.assertEqual(
            analysis.status,
            ResumeAnalysis.Status.PENDING,
        )
        self.assertIn(
            "Shifana Barveen",
            analysis.extracted_text,
        )

    def test_extract_and_save_resume_text_failure_persists_status(self):
        user = User.objects.create_user(
            email="failed_student@example.com",
            password="TestPassword123",
        )

        txt_file = SimpleUploadedFile(
            "resume.txt",
            b"This is an unsupported plain text file.",
            content_type="text/plain",
        )

        resume = Resume.objects.create(
            student=user,
            file=txt_file,
            original_filename="resume.txt",
            is_active=True,
        )

        with self.assertRaises(ValueError):
            extract_and_save_resume_text(resume)

        analysis = ResumeAnalysis.objects.get(resume=resume)
        self.assertEqual(
            analysis.status,
            ResumeAnalysis.Status.FAILED,
        )
        self.assertIn(
            "Unsupported resume format",
            analysis.error_message,
        )


@override_settings(GEMINI_API_KEY="test-gemini-key")
class GeminiServiceValidationTests(TestCase):

    def test_validate_resume_analysis_result_accepts_valid_payload(self):
        result = validate_resume_analysis_result(VALID_GEMINI_RESULT)

        self.assertEqual(result["resume_score"], 78)
        self.assertEqual(result["skills"], ["Python", "Java", "SQL"])

    def test_validate_resume_analysis_result_rejects_invalid_score(self):
        payload = {**VALID_GEMINI_RESULT, "resume_score": 150}

        with self.assertRaises(InvalidGeminiResponseError):
            validate_resume_analysis_result(payload)

    def test_validate_resume_analysis_result_rejects_missing_fields(self):
        payload = {"resume_score": 50}

        with self.assertRaises(InvalidGeminiResponseError):
            validate_resume_analysis_result(payload)

    def test_validate_resume_analysis_result_rejects_non_list_fields(self):
        payload = {**VALID_GEMINI_RESULT, "skills": "Python"}

        with self.assertRaises(InvalidGeminiResponseError):
            validate_resume_analysis_result(payload)


@override_settings(GEMINI_API_KEY="test-gemini-key")
class GeminiAnalysisServiceTests(TestCase):

    def create_docx(self):
        document = Document()
        document.add_paragraph("Shifana Barveen")
        document.add_paragraph("Python Java SQL")
        document.add_paragraph("Computer Science Engineering")

        docx_buffer = BytesIO()
        document.save(docx_buffer)
        docx_buffer.seek(0)

        return SimpleUploadedFile(
            "test_resume.docx",
            docx_buffer.read(),
            content_type=(
                "application/vnd.openxmlformats-officedocument"
                ".wordprocessingml.document"
            ),
        )

    @patch(
        "apps.resume_analysis.services.analysis_service.analyze_resume_with_gemini"
    )
    def test_successful_gemini_analysis_marks_completed(
        self,
        mock_analyze,
    ):
        mock_analyze.return_value = VALID_GEMINI_RESULT

        user = User.objects.create_user(
            email="gemini_student@example.com",
            password="TestPassword123",
        )
        resume = Resume.objects.create(
            student=user,
            file=self.create_docx(),
            original_filename="test_resume.docx",
            is_active=True,
        )

        analysis = analyze_resume(resume)

        self.assertEqual(
            analysis.status,
            ResumeAnalysis.Status.COMPLETED,
        )
        self.assertEqual(
            analysis.analysis_result,
            VALID_GEMINI_RESULT,
        )
        self.assertEqual(analysis.error_message, "")

    @patch(
        "apps.resume_analysis.services.analysis_service.analyze_resume_with_gemini"
    )
    def test_gemini_failure_preserves_extracted_text(
        self,
        mock_analyze,
    ):
        mock_analyze.side_effect = GeminiAPIError(
            "Failed to analyze resume with Gemini."
        )

        user = User.objects.create_user(
            email="gemini_fail@example.com",
            password="TestPassword123",
        )
        resume = Resume.objects.create(
            student=user,
            file=self.create_docx(),
            original_filename="test_resume.docx",
            is_active=True,
        )

        with self.assertRaises(GeminiAPIError):
            analyze_resume(resume)

        analysis = ResumeAnalysis.objects.get(resume=resume)
        self.assertEqual(
            analysis.status,
            ResumeAnalysis.Status.FAILED,
        )
        self.assertIn(
            "Shifana Barveen",
            analysis.extracted_text,
        )
        self.assertEqual(analysis.analysis_result, {})
        self.assertIn(
            "Failed to analyze resume with Gemini.",
            analysis.error_message,
        )

    @patch(
        "apps.resume_analysis.services.analysis_service.analyze_resume_with_gemini"
    )
    def test_invalid_gemini_payload_is_not_stored(
        self,
        mock_analyze,
    ):
        mock_analyze.side_effect = InvalidGeminiResponseError(
            "Resume analysis response failed validation."
        )

        user = User.objects.create_user(
            email="invalid_payload@example.com",
            password="TestPassword123",
        )
        resume = Resume.objects.create(
            student=user,
            file=self.create_docx(),
            original_filename="test_resume.docx",
            is_active=True,
        )

        with self.assertRaises(InvalidGeminiResponseError):
            analyze_resume(resume)

        analysis = ResumeAnalysis.objects.get(resume=resume)
        self.assertEqual(analysis.status, ResumeAnalysis.Status.FAILED)
        self.assertEqual(analysis.analysis_result, {})

    def test_empty_resume_text_is_handled(self):
        user = User.objects.create_user(
            email="empty_text@example.com",
            password="TestPassword123",
        )

        writer = PdfWriter()
        writer.add_blank_page(width=612, height=792)
        pdf_buffer = BytesIO()
        writer.write(pdf_buffer)
        pdf_buffer.seek(0)

        pdf_file = SimpleUploadedFile(
            "blank.pdf",
            pdf_buffer.read(),
            content_type="application/pdf",
        )

        resume = Resume.objects.create(
            student=user,
            file=pdf_file,
            original_filename="blank.pdf",
            is_active=True,
        )

        with self.assertRaises(InvalidGeminiResponseError):
            analyze_resume(resume)

        analysis = ResumeAnalysis.objects.get(resume=resume)
        self.assertEqual(analysis.status, ResumeAnalysis.Status.FAILED)
        self.assertIn(
            "empty or unreadable",
            analysis.error_message,
        )


@override_settings(GEMINI_API_KEY="test-gemini-key")
class ResumeAnalysisAPITests(APITestCase):

    def setUp(self):
        self.user = User.objects.create_user(
            email="student@example.com",
            password="TestPassword123",
            first_name="Test",
            last_name="Student",
        )

        self.other_user = User.objects.create_user(
            email="other@example.com",
            password="TestPassword123",
            first_name="Other",
            last_name="Student",
        )

        self.docx_file = self.create_docx()

        self.resume = Resume.objects.create(
            student=self.user,
            file=self.docx_file,
            original_filename="test_resume.docx",
            is_active=True,
        )

        self.other_resume = Resume.objects.create(
            student=self.other_user,
            file=self.create_docx(),
            original_filename="other_resume.docx",
            is_active=True,
        )

        self.url = reverse(
            "resume-analyze",
            kwargs={"resume_id": self.resume.id},
        )

        self.gemini_patcher = patch(
            "apps.resume_analysis.services.analysis_service.analyze_resume_with_gemini",
            return_value=VALID_GEMINI_RESULT,
        )
        self.mock_analyze = self.gemini_patcher.start()
        self.addCleanup(self.gemini_patcher.stop)

    def create_docx(self):
        document = Document()
        document.add_paragraph("Shifana Barveen")
        document.add_paragraph("Python Java SQL")
        document.add_paragraph("Computer Science Engineering")

        docx_buffer = BytesIO()
        document.save(docx_buffer)
        docx_buffer.seek(0)

        return SimpleUploadedFile(
            "test_resume.docx",
            docx_buffer.read(),
            content_type=(
                "application/vnd.openxmlformats-officedocument"
                ".wordprocessingml.document"
            ),
        )

    def test_authenticated_user_can_analyze_own_resume(self):
        """Authenticated user can analyze own resume."""
        self.client.force_authenticate(user=self.user)

        response = self.client.post(self.url)

        self.assertEqual(
            response.status_code,
            status.HTTP_200_OK,
        )
        self.assertEqual(
            response.data["resume_id"],
            self.resume.id,
        )
        self.assertEqual(
            response.data["original_filename"],
            "test_resume.docx",
        )
        self.assertEqual(
            response.data["status"],
            ResumeAnalysis.Status.COMPLETED,
        )
        self.assertIn(
            "Shifana Barveen",
            response.data["extracted_text"],
        )
        self.assertEqual(
            response.data["analysis_result"]["resume_score"],
            78,
        )
        self.assertEqual(
            response.data["error_message"],
            "",
        )
        self.assertIn("id", response.data)
        self.assertIn("analysis_id", response.data)
        self.assertIn("created_at", response.data)
        self.assertIn("updated_at", response.data)

    def test_unauthenticated_request_returns_401(self):
        """Unauthenticated request returns 401."""
        response = self.client.post(self.url)

        self.assertEqual(
            response.status_code,
            status.HTTP_401_UNAUTHORIZED,
        )

    def test_user_cannot_analyze_another_users_resume_and_receives_404(self):
        """User cannot analyze another user's resume and receives 404."""
        self.client.force_authenticate(user=self.other_user)

        response = self.client.post(self.url)

        self.assertEqual(
            response.status_code,
            status.HTTP_404_NOT_FOUND,
        )

    def test_resume_analysis_record_is_created(self):
        """ResumeAnalysis record is created."""
        self.client.force_authenticate(user=self.user)

        self.assertFalse(
            ResumeAnalysis.objects.filter(
                resume=self.resume
            ).exists()
        )

        response = self.client.post(self.url)

        self.assertEqual(
            response.status_code,
            status.HTTP_200_OK,
        )
        self.assertTrue(
            ResumeAnalysis.objects.filter(
                resume=self.resume
            ).exists()
        )

        analysis = ResumeAnalysis.objects.get(resume=self.resume)
        self.assertEqual(
            analysis.status,
            ResumeAnalysis.Status.COMPLETED,
        )
        self.assertEqual(
            analysis.id,
            response.data["id"],
        )

    def test_extracted_docx_text_is_saved_correctly(self):
        """Extracted DOCX text is saved correctly."""
        self.client.force_authenticate(user=self.user)

        response = self.client.post(self.url)

        self.assertEqual(
            response.status_code,
            status.HTTP_200_OK,
        )

        analysis = ResumeAnalysis.objects.get(resume=self.resume)
        self.assertIn("Shifana Barveen", analysis.extracted_text)
        self.assertIn("Python Java SQL", analysis.extracted_text)
        self.assertIn(
            "Computer Science Engineering",
            analysis.extracted_text,
        )
        self.assertEqual(
            analysis.extracted_text,
            response.data["extracted_text"],
        )

    def test_reanalyzing_same_resume_does_not_create_duplicate_records(self):
        """Re-analyzing updates the same ResumeAnalysis record."""
        self.client.force_authenticate(user=self.user)

        first_response = self.client.post(self.url)
        self.assertEqual(
            first_response.status_code,
            status.HTTP_200_OK,
        )
        self.assertEqual(
            ResumeAnalysis.objects.filter(
                resume=self.resume
            ).count(),
            1,
        )

        updated_result = {
            **VALID_GEMINI_RESULT,
            "resume_score": 82,
        }
        self.mock_analyze.return_value = updated_result

        second_response = self.client.post(self.url)
        self.assertEqual(
            second_response.status_code,
            status.HTTP_200_OK,
        )
        self.assertEqual(
            ResumeAnalysis.objects.filter(
                resume=self.resume
            ).count(),
            1,
        )
        self.assertEqual(
            first_response.data["id"],
            second_response.data["id"],
        )
        self.assertEqual(
            second_response.data["analysis_result"]["resume_score"],
            82,
        )

    def test_invalid_nonexistent_resume_returns_404(self):
        """Invalid/nonexistent resume returns 404."""
        self.client.force_authenticate(user=self.user)

        nonexistent_url = reverse(
            "resume-analyze",
            kwargs={"resume_id": 99999},
        )
        response = self.client.post(nonexistent_url)

        self.assertEqual(
            response.status_code,
            status.HTTP_404_NOT_FOUND,
        )

    def test_extraction_failure_is_handled_correctly(self):
        """Extraction failure is handled correctly."""
        txt_file = SimpleUploadedFile(
            "unsupported.txt",
            b"This is a plain text file.",
            content_type="text/plain",
        )
        unsupported_resume = Resume.objects.create(
            student=self.user,
            file=txt_file,
            original_filename="unsupported.txt",
            is_active=True,
        )

        self.client.force_authenticate(user=self.user)

        fail_url = reverse(
            "resume-analyze",
            kwargs={"resume_id": unsupported_resume.id},
        )
        response = self.client.post(fail_url)

        self.assertEqual(
            response.status_code,
            status.HTTP_400_BAD_REQUEST,
        )
        self.assertEqual(
            response.data["status"],
            ResumeAnalysis.Status.FAILED,
        )
        self.assertIn(
            "Unsupported resume format",
            response.data["error_message"],
        )

        analysis = ResumeAnalysis.objects.get(
            resume=unsupported_resume
        )
        self.assertEqual(
            analysis.status,
            ResumeAnalysis.Status.FAILED,
        )
        self.assertIn(
            "Unsupported resume format",
            analysis.error_message,
        )

    def test_gemini_failure_returns_failed_status(self):
        self.mock_analyze.side_effect = GeminiAPIError(
            "Failed to analyze resume with Gemini."
        )

        self.client.force_authenticate(user=self.user)
        response = self.client.post(self.url)

        self.assertEqual(
            response.status_code,
            status.HTTP_400_BAD_REQUEST,
        )
        self.assertEqual(
            response.data["status"],
            ResumeAnalysis.Status.FAILED,
        )
        self.assertIn(
            "Failed to analyze resume with Gemini.",
            response.data["error_message"],
        )

        analysis = ResumeAnalysis.objects.get(resume=self.resume)
        self.assertIn(
            "Shifana Barveen",
            analysis.extracted_text,
        )
        self.assertEqual(analysis.analysis_result, {})

    def test_get_analysis(self):
        self.client.force_authenticate(user=self.user)

        response = self.client.get(self.url)
        self.assertEqual(
            response.status_code,
            status.HTTP_404_NOT_FOUND,
        )

        self.client.post(self.url)
        response = self.client.get(self.url)
        self.assertEqual(
            response.status_code,
            status.HTTP_200_OK,
        )
        self.assertEqual(
            response.data["status"],
            ResumeAnalysis.Status.COMPLETED,
        )
        self.assertEqual(
            response.data["resume_id"],
            self.resume.id,
        )

    @patch(
        "apps.resume_analysis.services.gemini_service.genai.Client"
    )
    def test_analyze_resume_with_gemini_parses_json_response(
        self,
        mock_client,
    ):
        mock_response = mock_client.return_value.models.generate_content.return_value
        mock_response.text = json.dumps(VALID_GEMINI_RESULT)

        result = analyze_resume_with_gemini(
            "Shifana Barveen\nPython Java SQL"
        )

        self.assertEqual(result["resume_score"], 78)
        mock_client.return_value.models.generate_content.assert_called_once()

    @patch(
        "apps.resume_analysis.services.gemini_service.genai.Client"
    )
    def test_analyze_resume_with_gemini_rejects_invalid_json(
        self,
        mock_client,
    ):
        mock_response = mock_client.return_value.models.generate_content.return_value
        mock_response.text = "not-json"

        with self.assertRaises(InvalidGeminiResponseError):
            analyze_resume_with_gemini(
                "Shifana Barveen\nPython Java SQL"
            )
